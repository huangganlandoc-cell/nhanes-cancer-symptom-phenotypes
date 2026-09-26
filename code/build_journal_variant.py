"""Derive a journal-specific manuscript from the single source and build its submission folder.

The source is always `manuscript/Manuscript.md`, which is written in BMC style. A variant is produced
by a small set of declared transforms (abstract form, section names, declarations block, reference
style, how supplementary material is cited), so that no journal version is ever edited by hand and a
change to the source propagates everywhere with one command.

    python3 code/build_journal_variant.py bmc_public_health
    python3 code/build_journal_variant.py --list

Each variant writes submission/<folder>/manuscript_source.md (the derived markdown, kept for the
record) plus the Word files, figures and additional files, numbered in upload order.
"""
import argparse, datetime, json, re, shutil, subprocess, sys, tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))


def pandoc(md_text, out_path, title):
    """Markdown -> docx; east_asian_line_breaks keeps Chinese text from gaining stray spaces."""
    from docx import Document
    import build_docx
    src = Path("/tmp/_variant_src.md")
    src.write_text(md_text, encoding="utf-8")
    # no --metadata title: pandoc would print it as a visible first line ("Cover letter - ..."); the title
    # goes into the file properties instead. hard_line_breaks keeps the address and signature lines apart.
    subprocess.run(["pandoc", "-f", "markdown+east_asian_line_breaks+hard_line_breaks", "-t", "docx",
                    str(src), "-o", str(out_path)], check=True)
    d = Document(str(out_path))
    build_docx.set_properties(d, out_path)
    d.save(str(out_path))

def landscape(docx_path):
    """The STROBE checklist is five columns wide and is unreadable on a portrait page."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches
    d = Document(str(docx_path))
    for sec in d.sections:
        w, h = sec.page_width, sec.page_height          # pandoc leaves the page size unset
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = (h, w) if w and h else (Inches(11), Inches(8.5))
        sec.left_margin = sec.right_margin = Inches(0.6)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
    d.save(str(docx_path))


def to_pdf(docx_path, out_dir):
    """LibreOffice headless conversion; Nature Portfolio asks for the supplementary file as one PDF."""
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return Path(out_dir) / (Path(docx_path).stem + ".pdf")


SOURCE = "manuscript/Manuscript.md"
FIGS = ["figures/Figure1_symptom_phenotypes.png",
        "figures/Figure2_survival_curves.png",
        "figures/Figure3_sensitivity.png"]


# --------------------------------------------------------------------------- transforms
def strip_structured_abstract(md):
    """Turn the BMC structured abstract into one continuous paragraph (Nature-style journals)."""
    i = md.index("## Abstract\n\n") + len("## Abstract\n\n")
    j = md.index("\n**Keywords**")
    body = md[i:j]
    parts = re.findall(r"\*\*(?:Background|Methods|Results|Conclusions)\*\*\s*(.+?)(?=\n\n\*\*|\Z)", body, re.S)
    if not parts:
        return md
    return md[:i] + " ".join(re.sub(r"\s+", " ", p).strip() for p in parts) + md[j:]


def rename_abstract_heads(md, mapping):
    for old, new in mapping.items():
        md = md.replace(f"**{old}** ", f"**{new}** ", 1)
    return md


def rename_section(md, old, new):
    return md.replace(f"## {old}\n", f"## {new}\n")


def supplementary_wording(md, style):
    """BMC cites 'Additional file 1: Table S1'; other journals cite 'Supplementary Table S1'."""
    if style == "supplementary":
        md = re.sub(r"Additional file 1: (Tables?|Figs?\.) S", r"Supplementary \1 S", md)
        md = md.replace("**Additional file 1.**", "**Supplementary Material 1.**")
        md = md.replace("**Additional file 2.**", "**Supplementary Material 2.**")
        md = md.replace("(Additional file 2)", "(Supplementary Material 2)")
        assert "Additional file" not in md.partition("## Display items")[0], "an Additional file reference survived"
    if style == "wiley":
        # Wiley (Cancer Medicine): supporting files are Data S1 (figures and tables) and Data S2 (STROBE);
        # in the text items are cited as "Table S1" and "Figure S1" without a file prefix
        md = re.sub(r"Additional file 1: Tables S", "Tables S", md)
        md = re.sub(r"Additional file 1: Table S", "Table S", md)
        md = re.sub(r"Additional file 1: Figs?\. S", "Figure S", md)
        md = md.replace("(Additional file 2)", "(Data S2)")
        body, sep, manifest = md.partition("## Display items")
        body = re.sub(r"(?<![A-Za-z])Fig\. (S\d)", r"Figure \1", body)
        assert "Additional file" not in body, "an Additional file reference survived the Wiley transform"
        md = body + sep + manifest
    if style == "nature":
        # Nature Portfolio: one combined Supplementary Information file, and the word "Supplementary"
        # on every mention, including the bare "Table S11" forms the BMC source is free to use.
        md = re.sub(r"Additional file 1: (Tables?|Figs?\.) S", r"Supplementary \1 S", md)
        md = re.sub(r"(?<!Supplementary )\b(Tables?|Figs?\.|Figures?) (S\d)", r"Supplementary \1 \2", md)
        md = re.sub(r"Supplementary Table (S\d+) and (S\d+)", r"Supplementary Tables \1 and \2", md)
        body, sep, manifest = md.partition("## Display items")
        manifest = re.sub(r"^\|\s*Supplementary (Figure|Table) (S\d+)\s*\|", r"| \1 \2 |",
                          manifest, flags=re.M)              # the manifest keys must stay as build_docx wants
        body = body.replace("## Supplementary information", "## Supplementary Information")
        body = body.replace("**Additional file 1.** Supplementary figures S1 to S3 and supplementary "
                            "tables S1 to S15, with their legends.\n\n**Additional file 2.** STROBE "
                            "checklist for cohort studies.",
                            "Supplementary Information accompanies this paper: Supplementary Figs. S1 to "
                            "S3 and Supplementary Tables S1 to S15 with their legends, followed by the "
                            "STROBE checklist for cohort studies.")
        body = body.replace("(Additional file 2)", "(Supplementary Information)")
        md = body + sep + manifest
        assert "Additional file" not in md, "an Additional file reference survived the Nature transform"
    return md


def references(md, style, tmp):
    tmp.write_text(md, encoding="utf-8")
    subprocess.run([sys.executable, "code/build_references.py", style, "--apply", str(tmp)], check=True,
                   stdout=subprocess.DEVNULL)
    return tmp.read_text(encoding="utf-8")


# --------------------------------------------------------------------------- variants
VARIANTS = {
    "bmc_cancer": dict(numbering=True, folder="BMC_Cancer", journal="BMC Cancer", refs="bmc", abstract="structured",
                       supp="additional", notes="the source is already in this journal's style"),
    # BMC Public Health: requirements identical to BMC Cancer, but its recent NHANES papers cite
    # "Supplementary Table S1" rather than "Additional file 1: Table S1", so follow that practice
    "bmc_public_health": dict(numbering=True, folder="BMC_Public_Health", journal="BMC Public Health", refs="bmc",
                              abstract="structured", supp="supplementary",
                              notes="BMC series format; no public-database clause; supplementary wording"),
    # Scientific Reports: unstructured abstract of 200 words, Nature referencing, no abbreviations
    # list, ethics inside the Methods, and a 4,500-word main text (Methods excluded) that the
    # variant edits in manuscript/variant_edits/scientific_reports.json bring the text down to
    "scientific_reports": dict(folder="Scientific_Reports", journal="Scientific Reports", refs="nature",
                               abstract="override", supp="nature", drop_sections=["List of abbreviations"],
                               numbering=True, supp_pdf=True,
                               notes="Nature Portfolio format; main text capped at 4,500 words"),

    # Cancer Medicine (Wiley, open access): free-format submission (any consistent reference style); title
    # without abbreviations; 4 keywords (the guideline says both 1-4 and 4-6); structured abstract
    # Background/Methods/Results/Conclusion; Wiley end statements; supporting files Data S1 and Data S2.
    # The text is the Scientific Reports version (same shortening edits), so the two packages agree.
    "cancer_medicine": dict(numbering=True, folder="Cancer_Medicine", journal="Cancer Medicine", refs="bmc",
                            abstract="override", supp="wiley", edits_from="scientific_reports",
                            drop_sections=["Supplementary information", "List of abbreviations"],
                            title="Symptom Phenotypes and Mortality in United States Cancer Survivors: A Cohort Study",
                            keywords="cancer survivorship; latent class analysis; depressive symptoms; mortality",
                            supp_file="06_Data_S1_Supporting_Information.docx",
                            strobe_file="07_Data_S2_STROBE_checklist.docx",
                            notes="Wiley format; no word limit for original research; APC USD 4,970"),

    # Supportive Care in Cancer (Springer, hybrid): Purpose/Methods/Results/Conclusion abstract of
    # 150-250 words, Springer basic references, and a "Statements and Declarations" block that the
    # journal requires *after* the reference list
    "supportive_care": dict(numbering=True, folder="Supportive_Care_in_Cancer", journal="Supportive Care in Cancer",
                            refs="springer", abstract="structured", supp="supplementary",
                            declarations_after_references=True,
                            sections={"Introduction": "Introduction", "Conclusion": "Conclusion"},
                            notes="Springer format; body 3,500 words and 45 references are the stated caps"),
}


def build(key, out_root=Path("submission")):
    v = VARIANTS[key]
    md = Path(SOURCE).read_text(encoding="utf-8")
    out = out_root / v["folder"]
    out.mkdir(parents=True, exist_ok=True)
    override = Path("manuscript/abstracts") / f"{key}.md"
    if override.exists():          # a journal whose abstract rules force a different text of its own
        i = md.index("## Abstract\n\n") + len("## Abstract\n\n"); j = md.index("\n**Keywords**")
        md = md[:i] + override.read_text(encoding="utf-8").strip() + "\n" + md[j:]
    elif v["abstract"] == "unstructured":
        md = strip_structured_abstract(md)
    elif isinstance(v["abstract"], dict):
        md = rename_abstract_heads(md, v["abstract"])
    for old, new in v.get("sections", {}).items():
        md = rename_section(md, old, new)
    if v.get("title"):            # a journal-specific title (e.g. no abbreviations, Title Case)
        first, _, rest = md.partition("\n")
        assert first.startswith("# "), first
        md = f"# {v['title']}\n" + rest
    if v.get("keywords"):
        md, n = re.subn(r"^\*\*Keywords\*\*: .*$", "**Keywords**: " + v["keywords"], md, count=1, flags=re.M)
        assert n == 1, "no Keywords line"
    edits = Path("manuscript/variant_edits") / f"{v.get('edits_from', key)}.json"
    if edits.exists():             # sentences this journal's limits force us to shorten; source untouched
        for a, b in json.loads(edits.read_text(encoding="utf-8")):
            assert md.count(a) == 1, f"variant edit not unique ({md.count(a)}): {a[:60]!r}"
            md = md.replace(a, b)
    for head in v.get("drop_sections", []):
        i = md.index(f"## {head}")
        j = md.index("\n## ", i + 3)
        k = md.rfind("\n---\n", 0, i)          # drop the rule directly above the section, and nothing else
        start = k + 1 if k >= 0 and not md[k + 5:i].strip() else i
        md = md[:start] + md[j + 1:]
    md = supplementary_wording(md, v["supp"])
    decl = Path("manuscript/declarations") / f"{key}.md"
    tail = None
    if decl.exists():              # journals outside the BMC series want a different declarations block
        i = md.index("## Declarations"); j = md.index("\n## ", i + 5)
        block = decl.read_text(encoding="utf-8").strip()
        if v.get("declarations_after_references"):
            md = md[:i] + md[j + 1:]          # drop it from the body; it is rendered after the references
            tail = block
        else:
            md = md[:i] + block + "\n" + md[j:]
    src = out / "manuscript_source.md"
    src.write_text(md, encoding="utf-8")
    md = references(md, v["refs"], src)
    combined = v.get("supp_pdf")          # Nature Portfolio wants one supplementary file, preferably PDF
    title_text = md.split("\n", 1)[0].lstrip("#").strip()   # the variant's own title (e.g. no abbreviations)
    title = {"nature": "Supplementary Information", "supplementary": "Supplementary Material 1",
             "wiley": "Data S1. Supporting Information"}.get(v["supp"],
                                                                                              "Additional file 1")
    prefix = "Supplementary " if combined else ""
    if combined:
        intro = ('Supplementary Figures S1 to S3, Supplementary Tables S1 to S15 and the STROBE checklist for '
                 f'cohort studies, for "{title_text}".')
    elif v["supp"] == "wiley":
        intro = f'Figures S1–S3 and Tables S1–S15 for "{title_text}".'
    else:
        intro = f'Supplementary figures S1 to S3 and supplementary tables S1 to S15 for "{title_text}".'
    strobe = Path("supporting/STROBE_checklist_EN.md").read_text(encoding="utf-8")
    line = "Manuscript: Symptom phenotypes and mortality in US cancer survivors: a cohort study."
    assert line in strobe, "STROBE checklist title line changed"
    strobe = strobe.replace(line, f"Manuscript: {title_text}.")
    with tempfile.TemporaryDirectory() as tmp_dir:
        supp = Path(tmp_dir) / "Supplementary_Information.docx" if combined else out / v.get("supp_file", "06_Additional_file_1.docx")
        subprocess.run([sys.executable, "-c",
                        "import sys; sys.path.insert(0, 'code'); import build_docx; "
                        f"build_docx.build({str(src)!r}, {str(out / '02_Manuscript.docx')!r}, tail_md={tail!r}, "
                        f"numbering={bool(v.get('numbering'))}); "
                        f"build_docx.build_additional_file({str(src)!r}, {str(supp)!r}, title={title!r}, "
                        f"label_prefix={prefix!r}, intro={intro!r})"],
                       check=True, stdout=subprocess.DEVNULL)
        if combined:
            from pypdf import PdfWriter
            import fill_strobe_pages
            checklist = Path(tmp_dir) / "STROBE_checklist.docx"
            text = fill_strobe_pages.fill(strobe, out / "02_Manuscript.docx")   # needs the paginated file
            text = supplementary_wording(text, "nature")
            pandoc(text.replace("# STROBE checklist", "# Supplementary Information: STROBE checklist"),
                   checklist, "STROBE checklist")
            landscape(checklist)
            writer = PdfWriter()
            for f in (to_pdf(supp, tmp_dir), to_pdf(checklist, tmp_dir)):
                writer.append(str(f))
            writer.add_metadata({"/Title": "06_Supplementary_Information", "/Author": "Yongyong Bao"})
            writer.write(str(out / "06_Supplementary_Information.pdf")); writer.close()
        else:
            import fill_strobe_pages
            text = strobe
            if v.get("numbering"):             # page numbers exist only in a paginated manuscript
                text = fill_strobe_pages.fill(strobe, out / "02_Manuscript.docx", si="Data S1" if v["supp"] == "wiley" else title)
            text = supplementary_wording(text, v["supp"])
            label = {"supplementary": "Supplementary Material 2", "wiley": "Data S2"}.get(v["supp"], "Additional file 2")
            pandoc(text.replace("# STROBE checklist", f"# {label}. STROBE checklist"),
                   out / v.get("strobe_file", "07_Additional_file_2_STROBE_checklist.docx"), label)
    for i, f in enumerate(FIGS, start=3):
        shutil.copy(f, out / f"0{i}_Figure_{i - 2}.png")
    letter = Path("manuscript/cover_letters") / f"{key}.md"
    if letter.exists():
        body = letter.read_text(encoding="utf-8")
        assert "to be completed" not in body and "to be provided" not in body, f"{letter} still has placeholders"
        today = datetime.date.today()
        body = body.replace("[Date]", f"{today.day} {today:%B %Y}")   # dated when the package is built
        pandoc(body, out / "01_Cover_letter.docx", f"Cover letter - {v['journal']}")
    else:
        print(f"   (no cover letter yet: write {letter})")
    print(f"{v['journal']}: {out}/")
    for f in sorted(out.iterdir()):
        print(f"   {f.name:46s} {f.stat().st_size / 1024:7.0f} KB")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("variant", nargs="?", choices=sorted(VARIANTS))
    ap.add_argument("--list", action="store_true")
    a = ap.parse_args()
    if a.list or not a.variant:
        for k, v in VARIANTS.items():
            print(f"{k:20s} -> submission/{v['folder']}  ({v['journal']}; {v['notes']})")
    else:
        build(a.variant)
